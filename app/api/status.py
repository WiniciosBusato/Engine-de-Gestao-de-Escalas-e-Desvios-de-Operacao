import io
import csv
from fastapi.responses import StreamingResponse
from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session
from datetime import datetime, date
from typing import Optional
from app.core.database import get_db
from app.models import models
from app.schemas import schemas
from app.services.adherence import check_adherence, calculate_daily_adherence, get_agent_infractions, get_team_realtime_dashboard
from app.api.deps import require_roles
from app.models.models import UserRole, User


# A variável que o main.py está procurando:
router = APIRouter(prefix="/status", tags=["Status & Adherence"])

@router.post("/", response_model=schemas.StatusLogResponse)
def log_agent_status(payload: schemas.StatusLogCreate, db: Session = Depends(get_db)):
    # 1. Confere se o agente existe
    agent = db.query(models.Agent).filter(models.Agent.id == payload.agent_id).first()
    if not agent:
        raise HTTPException(status_code=404, detail="Agente não encontrado")

    now = payload.timestamp or datetime.now()

    # 2. Busca o último status aberto do agente para fechar a duração
    last_log = (
        db.query(models.StatusLog)
        .filter(models.StatusLog.agent_id == payload.agent_id, models.StatusLog.duration_seconds.is_(None))
        .order_by(models.StatusLog.timestamp.desc())
        .first()
    )

    if last_log:
        elapsed = int((now - last_log.timestamp).total_seconds())
        last_log.duration_seconds = max(elapsed, 0)
        db.add(last_log)

    # 3. Cria o novo registro de status
    new_log = models.StatusLog(
        agent_id=payload.agent_id,
        status=payload.status,
        timestamp=now
    )
    db.add(new_log)
    db.commit()
    db.refresh(new_log)

    return new_log

@router.get("/adherence/overview", response_model=schemas.AdherenceOverviewResponse)
def get_adherence_overview(
    check_time: Optional[datetime] = None,
    grace_period_minutes: int = 3,
    db: Session = Depends(get_db)
):
    agents = db.query(models.Agent).all()
    now = datetime.now()

    overview_items = []
    adherent_count = 0

    for agent in agents:
        #busca o ultimo status registrado para o operador
        latest_log = (
            db.query(models.StatusLog)
            .filter(models.StatusLog.agent_id == agent.id)
            .order_by(models.StatusLog.timestamp.desc())
            .first()
        )
        current_status = latest_log.status if latest_log else models.AgentStatus.OFFLINE

        #execute a regra de negocio de aderencia
        is_adherent, expected_status, message = check_adherence(
            db=db,
            agent_id=agent.id,
            current_status=current_status,
            check_time=now
        )

        if is_adherent:
            adherent_count += 1

        overview_items.append(
            schemas.AgentOverviewItem(
                agent_id=agent.id,
                agent_name=agent.name,
                skill_group=agent.skill_group,
                current_status=current_status,
                expected_status=expected_status,
                is_adherent=is_adherent,
                message=message
            )
        )

    total_agents = len(agents)
    rate = round((adherent_count / total_agents* 100), 2) if total_agents > 0 else 0.0

    return schemas.AdherenceOverviewResponse(
        total_agents=total_agents,
        adherent_count=adherent_count,
        non_adherent_count=total_agents - adherent_count,
        adherence_rate=rate,
        timestamp=now,
        agents=overview_items
    )


@router.get("/adherence/daily/{agent_id}", response_model=schemas.DailyAdherenceResponse)
def get_daily_adherence(
    agent_id: int,
    report_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(UserRole.SUPERVISOR, UserRole.ADMIN))
):
    agent = db.query(models.Agent).filter(models.Agent.id == agent_id).first()
    if not agent:
        raise HTTPException(status_code=404, detail="Agente nao encontrado")

    target_date = report_date or date.today()
    result = calculate_daily_adherence(db=db, agent_id=agent_id, target_date=target_date)

    if not result:
        raise HTTPException(status_code=404, detail=f"Escala nao encontrada para a data {target_date}")

    return {
        "agent_id": agent.id,
        "agent_name": agent.name,
        "date": target_date,
        "total_planned_seconds": result["total_planned_seconds"],
        "total_adherent_seconds": result["total_adherent_seconds"],
        "overall_adherence_rate": result["overall_adherence_rate"],
        "intervals": result.get("intervals", [])
    }

def get_daily_adherence(
    agent_id: int,
    report_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(UserRole.SUPERVISOR, UserRole.ADMIN))
):
    agent = db.query(models.Agent).filter(models.Agent.id == agent_id).first()
    if not agent:
        raise HTTPException(status_code=404, detail="Agente nao encontrado")

    target_date = report_date or date.today()
    result = calculate_daily_adherence(db=db, agent_id=agent_id, target_date=target_date)

    if not result:
        raise HTTPException(status_code=404, detail=f"Escala nao encontrada para a data {target_date}")

    return {
        "agent_id": agent.id,
        "agent_name": agent.name,
        "date": target_date,
        "total_planned_seconds": result["total_planned_seconds"],
        "total_adherent_seconds": result["total_adherent_seconds"],
        "overall_adherence_rate": result["overall_adherence_rate"],
        "intervals": result.get("intervals", [])
    }
@router.get("/adherence/{agent_id}", response_model=schemas.AdherenceCheckResponse)
def get_agent_adherence(
    agent_id: int,
    check_time: Optional[datetime] = None,
    grace_period_minutes: int = 3,
    db: Session = Depends(get_db)
):
    agent = db.query(models.Agent).filter(models.Agent.id == agent_id).first()
    if not agent:
        raise HTTPException(status_code=404, detail="Agente não encontrado")

    result = check_adherence(
        db=db,
        agent_id=agent_id,
        check_time=check_time,
        grace_period_minutes=grace_period_minutes
    )
    status_msg = "Agente adrente a escala" if result.get("is_adherent") else "Agente fora de aderencia"

    return {
        **result,
        "checked_at": result.get("check_time") or datetime.now(),
        "message": status_msg
    }


@router.get("/adherence/{agent_id}/infractions", response_model=schemas.AgentInfractionResponse)
def get_agent_infractions_report(
    agent_id: int,
    target_date: Optional[date] = None,
    grace_period_minutes: int = 0,
    db: Session = Depends(get_db)
):
    query_date = target_date or date.today()
    result = get_agent_infractions(db=db, agent_id=agent_id, target_date=query_date, grace_period_minutes=grace_period_minutes)

    if not result:
        raise HTTPException(
            status_code=404,
            detail=f"Nenhuma escala encontrada para o agente {agent_id} na data {query_date}."
        )
    return result

@router.get("/adherence/realtime/team", response_model=schemas.TeamRealtimeDashboardResponse)
def get_realtime_team_adherence(
    grace_period_minutes: int = 0,
    db: Session = Depends(get_db)
):
    """
    Retorna o painel consolidado em tempo real de toda a equipe,
    com status atual, aderencia e duração de cada operador.
    """
    return get_team_realtime_dashboard(
        db=db,
        grace_period_minutes=grace_period_minutes
    )


@router.get("/adherence/daily/{agent_id}/export/csv")
def export_daily_adherence_csv(
    agent_id: int,
    report_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(UserRole.SUPERVISOR, UserRole.ADMIN))
):
    agent = db.query(models.Agent).filter(models.Agent.id == agent_id).first()
    if not agent:
        raise HTTPException(status_code=404, detail="Agente nao encontrado")

    target_date = report_date or date.today()
    result = calculate_daily_adherence(db=db, agent_id=agent_id, target_date=target_date)

    if not result:
        raise HTTPException(status_code=404, detail=f"Escala nao encontrada para a data {target_date}")

    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")

    # Cabecalho do Relatorio Consolidado
    writer.writerow(["ID Agente", "Nome Agente", "Data", "Segundos Planejados", "Segundos Aderentes", "Taxa Aderencia (%)"])
    writer.writerow([
        agent.id,
        agent.name,
        target_date.isoformat(),
        result["total_planned_seconds"],
        result["total_adherent_seconds"],
        result["overall_adherence_rate"]
    ])
    writer.writerow([])  # Linha em branco

    # Detalhamento de Intervalos
    writer.writerow(["Tipo Intervalo", "Inicio Planejado", "Fim Planejado", "Segundos Planejados", "Segundos Aderentes", "Aderencia (%)"])
    for interval in result.get("intervals", []):
        writer.writerow([
            interval.get("interval_type"),
            interval.get("planned_start"),
            interval.get("planned_end"),
            interval.get("planned_seconds"),
            interval.get("adherent_seconds"),
            interval.get("adherence_rate")
        ])

    output.seek(0)
    filename = f"aderencia_agente_{agent_id}_{target_date.isoformat()}.csv"

    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/adherence/daily/{agent_id}/export/excel")
def export_daily_adherence_excel(
    agent_id: int,
    report_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(UserRole.SUPERVISOR, UserRole.ADMIN))
):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    agent = db.query(models.Agent).filter(models.Agent.id == agent_id).first()
    if not agent:
        raise HTTPException(status_code=404, detail="Agente nao encontrado")

    target_date = report_date or date.today()
    result = calculate_daily_adherence(db=db, agent_id=agent_id, target_date=target_date)

    if not result:
        raise HTTPException(status_code=404, detail=f"Escala nao encontrada para a data {target_date}")

    wb = openpyxl.Workbook()
    
    # Aba 1: Resumo Executivo
    ws_resumo = wb.active
    ws_resumo.title = "Resumo Geral"

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    headers_resumo = ["ID Agente", "Nome Agente", "Data", "Segundos Planejados", "Segundos Aderentes", "Taxa Aderência (%)"]
    ws_resumo.append(headers_resumo)

    for col_num in range(1, len(headers_resumo) + 1):
        cell = ws_resumo.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws_resumo.append([
        agent.id,
        agent.name,
        target_date.isoformat(),
        result["total_planned_seconds"],
        result["total_adherent_seconds"],
        result["overall_adherence_rate"]
    ])

    for col_num in range(1, len(headers_resumo) + 1):
        cell = ws_resumo.cell(row=2, column=col_num)
        cell.font = data_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Aba 2: Detalhamento por Intervalo
    ws_detalhe = wb.create_sheet(title="Intervalos")
    headers_detalhe = ["Tipo Intervalo", "Início Planejado", "Fim Planejado", "Segundos Planejados", "Segundos Aderentes", "Aderência (%)"]
    ws_detalhe.append(headers_detalhe)

    for col_num in range(1, len(headers_detalhe) + 1):
        cell = ws_detalhe.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, interval in enumerate(result.get("intervals", []), start=2):
        ws_detalhe.append([
            interval.get("interval_type"),
            str(interval.get("planned_start")),
            str(interval.get("planned_end")),
            interval.get("planned_seconds"),
            interval.get("adherent_seconds"),
            interval.get("adherence_rate")
        ])
        for col_num in range(1, len(headers_detalhe) + 1):
            cell = ws_detalhe.cell(row=idx, column=col_num)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Auto-ajuste de largura de coluna nas duas abas
    for sheet in [ws_resumo, ws_detalhe]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    filename = f"aderencia_agente_{agent_id}_{target_date.isoformat()}.xlsx"

    return StreamingResponse(
        excel_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
